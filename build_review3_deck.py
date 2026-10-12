import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_review3_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Human Academic Palette
    C_WHITE        = RGBColor(255, 255, 255)
    C_BG_PAGE      = RGBColor(248, 250, 252)   # #F8FAFC
    C_NAVY_TITLE   = RGBColor(30, 58, 138)     # #1E3A8A
    C_NAVY_DARK    = RGBColor(15, 23, 42)      # #0F172A
    C_BLUE_LINE    = RGBColor(37, 99, 235)     # #2563EB
    C_TEXT_MAIN    = RGBColor(30, 41, 59)      # #1E293B
    C_TEXT_MUTED   = RGBColor(100, 116, 139)   # #64748B
    
    # Soft box colors
    C_BOX_BORDER   = RGBColor(203, 213, 225)   # #CBD5E1
    C_BLUE_BG      = RGBColor(239, 246, 255)   # #EFF6FF
    C_BLUE_BORDER  = RGBColor(147, 197, 253)   # #93C5FD
    C_GREEN_BG     = RGBColor(240, 253, 244)   # #F0FDF4
    C_GREEN_BORDER = RGBColor(134, 239, 172)   # #86EFAC
    C_GREEN_TEXT   = RGBColor(21, 128, 61)     # #15803D
    C_AMBER_BG     = RGBColor(254, 243, 199)   # #FEF3C7
    C_AMBER_BORDER = RGBColor(252, 211, 77)    # #FCD34D
    C_AMBER_TEXT   = RGBColor(180, 83, 9)      # #B45309
    C_RED_BG       = RGBColor(254, 242, 242)   # #FEF2F2
    C_RED_BORDER   = RGBColor(252, 165, 165)   # #FCA5A5
    C_RED_TEXT     = RGBColor(185, 28, 28)     # #B91C1C
    C_PURPLE_BG    = RGBColor(245, 243, 255)   # #F5F3FF
    C_PURPLE_BORDER= RGBColor(196, 181, 253)   # #C4B5FD
    C_PURPLE_TEXT  = RGBColor(109, 40, 217)    # #6D28D9

    TOTAL_SLIDES = 11

    def add_slide_base():
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG_PAGE
        bg.line.fill.background()
        return slide

    def add_slide_header(slide, title_text, slide_num):
        t_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.45), Inches(10.5), Inches(0.75))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_top = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Calibri"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_TITLE

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.15), Inches(2.2), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = C_BLUE_LINE
        line.line.fill.background()

        n_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.45), Inches(1.0), Inches(0.4))
        tf_n = n_box.text_frame
        p_n = tf_n.paragraphs[0]
        p_n.text = f"{slide_num} / {TOTAL_SLIDES}"
        p_n.font.name = "Calibri"
        p_n.font.size = Pt(11)
        p_n.font.color.rgb = C_TEXT_MUTED
        p_n.alignment = PP_ALIGN.RIGHT

        f_box = slide.shapes.add_textbox(Inches(0.9), Inches(7.1), Inches(10.0), Inches(0.3))
        tf_f = f_box.text_frame
        p_f = tf_f.paragraphs[0]
        p_f.text = "Academic Conference DBMS  •  Review 3 (Final Review: Web Application & System Demo)"
        p_f.font.name = "Calibri"
        p_f.font.size = Pt(9.5)
        p_f.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    slide1 = add_slide_base()
    
    v_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.4), Inches(7.5))
    v_bar.fill.solid()
    v_bar.fill.fore_color.rgb = C_NAVY_TITLE
    v_bar.line.fill.background()

    t_box1 = slide1.shapes.add_textbox(Inches(1.5), Inches(1.3), Inches(10.5), Inches(3.2))
    tf1 = t_box1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "DATABASE MANAGEMENT SYSTEMS  |  CS304"
    p0.font.name = "Calibri"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = C_BLUE_LINE
    p0.space_after = Pt(12)

    p1 = tf1.add_paragraph()
    p1.text = "Academic Conference Paper Submission\n& Review System DBMS"
    p1.font.name = "Calibri"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = C_NAVY_DARK
    p1.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "Review 3 (Final Review): Front-End Web Application & End-to-End Demonstration"
    p2.font.name = "Calibri"
    p2.font.size = Pt(16)
    p2.font.color.rgb = C_TEXT_MUTED

    meta_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.7), Inches(10.2), Inches(1.8))
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = C_WHITE
    meta_box.line.color.rgb = C_BOX_BORDER
    meta_box.line.width = Pt(1)

    tf_m = meta_box.text_frame
    tf_m.margin_left = Inches(0.4)
    tf_m.margin_top = Inches(0.25)
    
    p = tf_m.paragraphs[0]
    p.text = "FINAL REVIEW PROJECT DETAILS"
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_TITLE
    p.space_after = Pt(8)

    p = tf_m.add_paragraph()
    p.text = "• Presenter / Team: [Your Name / Team Members]"
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = C_TEXT_MAIN
    p.space_after = Pt(4)

    p = tf_m.add_paragraph()
    p.text = "• Register / Roll No: [Your Register Number]"
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = C_TEXT_MAIN
    p.space_after = Pt(4)

    p = tf_m.add_paragraph()
    p.text = "• Full-Stack Architecture: Python Flask Web Engine  •  SQLite / PostgreSQL  •  Bootstrap 5 UI  •  Automated Tests"
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 2: REVIEW 3 OBJECTIVES & ARCHITECTURE
    # ==========================================
    slide2 = add_slide_base()
    add_slide_header(slide2, "Review 3 Objectives & Full-Stack System Architecture", 2)

    pillars = [
        ("1. Front-End Web Application", "Developed in Python Flask with Bootstrap 5 UI, providing responsive user portals for Authors, Reviewers, and Program Chairs.", C_BLUE_LINE),
        ("2. End-to-End Lifecycle", "Covers paper submission, version tracking, reviewer assignment, review entry, editorial decisions, and session scheduling.", C_PURPLE_TEXT),
        ("3. Business Rules in UI", "Enforces Conflict of Interest (COI) barring, 3-review quotas, and author presentation clash prevention in real time.", C_AMBER_TEXT),
        ("4. Systematic Testing", "Automated test suite (test_suite.py) verifying input constraints, triggers, and transaction rollbacks with 100% pass rate.", C_GREEN_TEXT)
    ]

    card_w = Inches(5.6)
    card_h = Inches(2.25)
    coords = [
        (Inches(0.9), Inches(1.6)),
        (Inches(6.8), Inches(1.6)),
        (Inches(0.9), Inches(4.15)),
        (Inches(6.8), Inches(4.15))
    ]

    for (title, desc, col), (x, y) in zip(pillars, coords):
        c = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = C_WHITE
        c.line.color.rgb = C_BOX_BORDER
        c.line.width = Pt(1)

        st = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(0.12), card_h)
        st.fill.solid()
        st.fill.fore_color.rgb = col
        st.line.fill.background()

        tf = c.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Calibri"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Calibri"
        p.font.size = Pt(12.5)
        p.font.color.rgb = C_TEXT_MAIN

    callout2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(6.55), Inches(11.5), Inches(0.42))
    callout2.fill.solid()
    callout2.fill.fore_color.rgb = C_BLUE_BG
    callout2.line.color.rgb = C_BLUE_BORDER
    tf_c2 = callout2.text_frame
    tf_c2.margin_top = Inches(0.06)
    p = tf_c2.paragraphs[0]
    p.text = "⚡ System Architecture: Browser Client  ⟶  Flask WSGI Controller  ⟶  Integrity Triggers  ⟶  Relational RDBMS"
    p.font.name = "Calibri"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_TITLE
    p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 3: ROLE-BASED PORTALS OVERVIEW
    # ==========================================
    slide3 = add_slide_base()
    add_slide_header(slide3, "Role-Based Front-End Portals Overview", 3)

    portals = [
        ("📝 Author Portal", "/author", "• Paper submission & abstract\n• Track selection & co-authors\n• Upload revision & camera-ready\n• View decision & feedback", C_BLUE_LINE),
        ("⚖ Reviewer Portal", "/reviewer", "• Blinded manuscript evaluation\n• Double-blind anonymity\n• 4 score criteria (1 to 10)\n• Separate author & chair comments", C_PURPLE_TEXT),
        ("👑 Program Chair", "/chair", "• COI-aware reviewer assignment\n• Quota-gated decision logging\n• Accept / Reject / Revise verdicts\n• Live committee audit log", C_AMBER_TEXT),
        ("📅 Timetable & Timings", "/schedule", "• Session allocation & hall booking\n• Author overlap collision check\n• Chronological presentation view\n• Master conference schedule", C_GREEN_TEXT)
    ]

    pw = Inches(2.72)
    ph = Inches(4.7)
    py = Inches(1.6)

    for idx, (pname, proute, pdesc, pcol) in enumerate(portals):
        px = Inches(0.9 + idx * 2.92)

        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, py, pw, ph)
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BOX_BORDER
        card.line.width = Pt(1)

        bhdr = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, py, pw, Inches(0.5))
        bhdr.fill.solid()
        bhdr.fill.fore_color.rgb = pcol
        bhdr.line.fill.background()
        tf_h = bhdr.text_frame
        tf_h.margin_top = Inches(0.1)
        p = tf_h.paragraphs[0]
        p.text = pname
        p.font.name = "Calibri"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER

        tf = card.text_frame
        tf.margin_left = Inches(0.15)
        tf.margin_top = Inches(0.65)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"Route: `{proute}`"
        p.font.name = "Consolas"
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_TITLE
        p.space_after = Pt(10)

        for line in pdesc.split("\n"):
            p = tf.add_paragraph()
            p.text = line
            p.font.name = "Calibri"
            p.font.size = Pt(11)
            p.font.color.rgb = C_TEXT_MAIN
            p.space_after = Pt(6)

    callout3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.42))
    callout3.fill.solid()
    callout3.fill.fore_color.rgb = C_BG_PAGE
    callout3.line.color.rgb = C_BOX_BORDER
    tf_c3 = callout3.text_frame
    tf_c3.margin_top = Inches(0.06)
    p = tf_c3.paragraphs[0]
    p.text = "✔ Multi-Role Experience: Each portal tailors UI visibility to respect conference peer review protocols."
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.color.rgb = C_TEXT_MUTED
    p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 4: MODULE 1: SUBMISSION & VERSIONS
    # ==========================================
    slide4 = add_slide_base()
    add_slide_header(slide4, "Module 1: Paper Submission & Multi-Version Lifecycle", 4)

    col_w = Inches(5.6)
    col_h = Inches(4.7)

    # Left: Submission Workflow
    b_sub = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.6), col_w, col_h)
    b_sub.fill.solid()
    b_sub.fill.fore_color.rgb = C_WHITE
    b_sub.line.color.rgb = C_BLUE_BORDER
    b_sub.line.width = Pt(1)

    tf_s = b_sub.text_frame
    tf_s.margin_left = Inches(0.35)
    tf_s.margin_top = Inches(0.3)
    p = tf_s.paragraphs[0]
    p.text = "MANUSCRIPT SUBMISSION WORKFLOW"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_TITLE
    p.space_after = Pt(12)

    sub_steps = [
        ("Input Validation", "Author selects Track, inputs Title, Abstract, and Co-authors"),
        ("Atomic Registration", "Creates row in `Papers`, establishes `Paper_Authors` link"),
        ("Automatic Version 1", "Generates initial blinded draft (`p{id}_v1_blind.pdf`)"),
        ("State Transition", "Status automatically initialized to `'SUBMITTED'`")
    ]
    for s_title, s_desc in sub_steps:
        p = tf_s.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = "✔ " + s_title + ":\n"
        r1.font.bold = True
        r1.font.size = Pt(12.5)
        r1.font.color.rgb = C_BLUE_LINE
        r2 = p.add_run()
        r2.text = "   " + s_desc
        r2.font.size = Pt(11.5)
        r2.font.color.rgb = C_TEXT_MAIN

    # Right: Version Tracking
    b_ver = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), col_w, col_h)
    b_ver.fill.solid()
    b_ver.fill.fore_color.rgb = C_WHITE
    b_ver.line.color.rgb = C_GREEN_BORDER
    b_ver.line.width = Pt(1)

    tf_v = b_ver.text_frame
    tf_v.margin_left = Inches(0.35)
    tf_v.margin_top = Inches(0.3)
    p = tf_v.paragraphs[0]
    p.text = "MULTI-VERSION TRACKING & AUDIT"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_GREEN_TEXT
    p.space_after = Pt(12)

    ver_points = [
        ("Iterative Revisions", "Uploads V2, V3 revisions with mandatory change logs"),
        ("Camera-Ready Flag", "`is_camera_ready = 1` uploaded only upon acceptance"),
        ("Immutable History", "Previous PDF drafts are retained; zero version overwrites"),
        ("Author De-anonymization", "Camera-ready version includes author names & affiliations")
    ]
    for v_title, v_desc in ver_points:
        p = tf_v.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = "✔ " + v_title + ":\n"
        r1.font.bold = True
        r1.font.size = Pt(12.5)
        r1.font.color.rgb = C_GREEN_TEXT
        r2 = p.add_run()
        r2.text = "   " + v_desc
        r2.font.size = Pt(11.5)
        r2.font.color.rgb = C_TEXT_MAIN

    # ==========================================
    # SLIDE 5: MODULE 2: REVIEWER ASSIGNMENT & COI
    # ==========================================
    slide5 = add_slide_base()
    add_slide_header(slide5, "Module 2: Reviewer Assignment & Real-Time COI Protection", 5)

    # 2 Comparison Cards: Happy Path vs Conflict Blocked
    b_hp = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.6), col_w, col_h)
    b_hp.fill.solid()
    b_hp.fill.fore_color.rgb = C_WHITE
    b_hp.line.color.rgb = C_GREEN_BORDER
    b_hp.line.width = Pt(1)

    tf_hp = b_hp.text_frame
    tf_hp.margin_left = Inches(0.35)
    tf_hp.margin_top = Inches(0.3)
    p = tf_hp.paragraphs[0]
    p.text = "VALID ASSIGNMENT (HAPPY PATH)"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_GREEN_TEXT
    p.space_after = Pt(12)

    hp_points = [
        ("Zero COI Verification", "Queries `Reviewer_Conflicts` for all co-authors of the paper"),
        ("Workload Capacity Check", "Ensures `current_assignments < max_workload`"),
        ("Database Commit", "Inserts into `Paper_Assignments` with due date"),
        ("Status Update", "Paper transitions from `'SUBMITTED'` to `'UNDER_REVIEW'`"),
        ("UI Feedback", "Displays success toast with assignment ID and deadline")
    ]
    for h_title, h_desc in hp_points:
        p = tf_hp.add_paragraph()
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = "✔ " + h_title + ":\n"
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = C_GREEN_TEXT
        r2 = p.add_run()
        r2.text = "   " + h_desc
        r2.font.size = Pt(11)
        r2.font.color.rgb = C_TEXT_MAIN

    # Right: COI Blocked
    b_blk = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), col_w, col_h)
    b_blk.fill.solid()
    b_blk.fill.fore_color.rgb = C_WHITE
    b_blk.line.color.rgb = C_RED_BORDER
    b_blk.line.width = Pt(1)

    tf_bk = b_blk.text_frame
    tf_bk.margin_left = Inches(0.35)
    tf_bk.margin_top = Inches(0.3)
    p = tf_bk.paragraphs[0]
    p.text = "COI VIOLATION BLOCKED (REAL-TIME)"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_RED_TEXT
    p.space_after = Pt(12)

    bk_points = [
        ("Simulated Scenario", "Assigning Reviewer 1 (Lamport) to Paper 1 (Author: Hinton)"),
        ("Conflict Detected", "Active `SAME_INSTITUTION` conflict identified"),
        ("Real-Time Interception", "Application aborts assignment BEFORE SQL commit"),
        ("Error Alert Dispatched", "UI displays danger alert detailing conflict type & author"),
        ("Zero Database Corruption", "Preserves academic peer review integrity and fairness")
    ]
    for b_title, b_desc in bk_points:
        p = tf_bk.add_paragraph()
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = "❌ " + b_title + ":\n"
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = C_RED_TEXT
        r2 = p.add_run()
        r2.text = "   " + b_desc
        r2.font.size = Pt(11)
        r2.font.color.rgb = C_TEXT_MAIN

    # ==========================================
    # SLIDE 6: MODULE 3: PEER REVIEW ENTRY
    # ==========================================
    slide6 = add_slide_base()
    add_slide_header(slide6, "Module 3: Peer Review Entry & Feedback Isolation", 6)

    rev_features = [
        ("Double-Blind Anonymity", "Reviewer sees only Paper Title and Abstract; author identities, emails, and institutions are masked.", C_BLUE_LINE),
        ("Multi-Criteria Ratings (1 to 10)", "Numeric scoring for Originality, Technical Soundness, Writing Clarity, and Overall Recommendation.", C_PURPLE_TEXT),
        ("Dual-Channel Feedback", "Public comments visible to author upon decision; confidential remarks strictly reserved for Program Chair.", C_AMBER_TEXT),
        ("Atomic Status Progression", "Inserting review updates `Paper_Assignments.status` to `'COMPLETED'`, unblocking chair decisions.", C_GREEN_TEXT)
    ]

    for (title, desc, col), (x, y) in zip(rev_features, coords):
        c = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = C_WHITE
        c.line.color.rgb = C_BOX_BORDER
        c.line.width = Pt(1)

        st = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(0.12), card_h)
        st.fill.solid()
        st.fill.fore_color.rgb = col
        st.line.fill.background()

        tf = c.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Calibri"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Calibri"
        p.font.size = Pt(12.5)
        p.font.color.rgb = C_TEXT_MAIN

    callout6 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(6.55), Inches(11.5), Inches(0.42))
    callout6.fill.solid()
    callout6.fill.fore_color.rgb = C_PURPLE_BG
    callout6.line.color.rgb = C_PURPLE_BORDER
    tf_c6 = callout6.text_frame
    tf_c6.margin_top = Inches(0.06)
    p = tf_c6.paragraphs[0]
    p.text = "✔ Feedback Isolation: Confidential chair comments are never exposed to authors in queries or UI views."
    p.font.name = "Calibri"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE_TEXT
    p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 7: MODULE 4: EDITORIAL DECISIONS
    # ==========================================
    slide7 = add_slide_base()
    add_slide_header(slide7, "Module 4: Editorial Decisions & Quota Gating", 7)

    # Left: Quota Enforcement
    b_qgate = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.6), col_w, col_h)
    b_qgate.fill.solid()
    b_qgate.fill.fore_color.rgb = C_WHITE
    b_qgate.line.color.rgb = C_AMBER_BORDER
    b_qgate.line.width = Pt(1)

    tf_qg = b_qgate.text_frame
    tf_qg.margin_left = Inches(0.35)
    tf_qg.margin_top = Inches(0.3)
    p = tf_qg.paragraphs[0]
    p.text = "REVIEW QUOTA GATING (≥ 3 REVIEWS)"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_AMBER_TEXT
    p.space_after = Pt(12)

    qg_points = [
        ("Business Rule", "Papers must receive at least 3 completed peer reviews"),
        ("Live Simulation", "Paper #6 has only 1 completed review and 2 pending"),
        ("UI Action Attempt", "Chair attempts to log decision for Paper #6"),
        ("Interception", "Application checks review count and blocks decision"),
        ("Result", "Paper remains in `'UNDER_REVIEW'` until quota is satisfied")
    ]
    for q_title, q_desc in qg_points:
        p = tf_qg.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = "⚡ " + q_title + ":\n"
        r1.font.bold = True
        r1.font.size = Pt(12.5)
        r1.font.color.rgb = C_AMBER_TEXT
        r2 = p.add_run()
        r2.text = "   " + q_desc
        r2.font.size = Pt(11.5)
        r2.font.color.rgb = C_TEXT_MAIN

    # Right: Verdict Recording
    b_verd = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), col_w, col_h)
    b_verd.fill.solid()
    b_verd.fill.fore_color.rgb = C_WHITE
    b_verd.line.color.rgb = C_BLUE_BORDER
    b_verd.line.width = Pt(1)

    tf_vd = b_verd.text_frame
    tf_vd.margin_left = Inches(0.35)
    tf_vd.margin_top = Inches(0.3)
    p = tf_vd.paragraphs[0]
    p.text = "VERDICT LOGGING & STATE TRANSITIONS"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_TITLE
    p.space_after = Pt(12)

    vd_points = [
        ("Acceptance Path", "`ACCEPT` logs decision and updates `Papers.status = 'ACCEPTED'`"),
        ("Camera-Ready Unlocked", "Enables author revision portal to upload final camera-ready PDF"),
        ("Rejection Path", "`REJECT` logs decision and marks paper as `'REJECTED'`"),
        ("Audit Trail", "Stores Chair ID, timestamp, and Program Committee feedback summary")
    ]
    for v_title, v_desc in vd_points:
        p = tf_vd.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = "✔ " + v_title + ":\n"
        r1.font.bold = True
        r1.font.size = Pt(12.5)
        r1.font.color.rgb = C_BLUE_LINE
        r2 = p.add_run()
        r2.text = "   " + v_desc
        r2.font.size = Pt(11.5)
        r2.font.color.rgb = C_TEXT_MAIN

    # ==========================================
    # SLIDE 8: MODULE 5: REGISTRATION & TIMETABLE
    # ==========================================
    slide8 = add_slide_base()
    add_slide_header(slide8, "Module 5: Author Registration & Collision-Free Timetable", 8)

    # Left: Registration Clearance
    b_reg = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.6), col_w, col_h)
    b_reg.fill.solid()
    b_reg.fill.fore_color.rgb = C_WHITE
    b_reg.line.color.rgb = C_GREEN_BORDER
    b_reg.line.width = Pt(1)

    tf_rg = b_reg.text_frame
    tf_rg.margin_left = Inches(0.35)
    tf_rg.margin_top = Inches(0.3)
    p = tf_rg.paragraphs[0]
    p.text = "AUTHOR REGISTRATION AUDIT"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_GREEN_TEXT
    p.space_after = Pt(12)

    rg_points = [
        ("Publication Prerequisite", "At least one author must register with fee payment ($650)"),
        ("Payment Status Check", "`Registrations.payment_status = 'PAID'` audited in views"),
        ("Transaction Tracking", "Stores unique transaction reference (e.g. TXN_IEEE_2026_01)"),
        ("Clearance Clearance", "Unlocks scheduling into conference presentation sessions")
    ]
    for r_title, r_desc in rg_points:
        p = tf_rg.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = "✔ " + r_title + ":\n"
        r1.font.bold = True
        r1.font.size = Pt(12.5)
        r1.font.color.rgb = C_GREEN_TEXT
        r2 = p.add_run()
        r2.text = "   " + r_desc
        r2.font.size = Pt(11.5)
        r2.font.color.rgb = C_TEXT_MAIN

    # Right: Schedule Collision
    b_sch = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), col_w, col_h)
    b_sch.fill.solid()
    b_sch.fill.fore_color.rgb = C_WHITE
    b_sch.line.color.rgb = C_BLUE_BORDER
    b_sch.line.width = Pt(1)

    tf_sc = b_sch.text_frame
    tf_sc.margin_left = Inches(0.35)
    tf_sc.margin_top = Inches(0.3)
    p = tf_sc.paragraphs[0]
    p.text = "COLLISION-FREE TIMETABLE GENERATION"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_TITLE
    p.space_after = Pt(12)

    sc_points = [
        ("Session Mapping", "Assigns accepted papers to Halls (A, B, C) and Sessions"),
        ("Author Clash Check", "Verifies that co-authors are NOT scheduled in simultaneous slots"),
        ("Room Collision Check", "Verifies room is free during requested start and end times"),
        ("Public Timetable View", "`vw_programme_schedule` displays chronological schedule")
    ]
    for s_title, s_desc in sc_points:
        p = tf_sc.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = "✔ " + s_title + ":\n"
        r1.font.bold = True
        r1.font.size = Pt(12.5)
        r1.font.color.rgb = C_BLUE_LINE
        r2 = p.add_run()
        r2.text = "   " + s_desc
        r2.font.size = Pt(11.5)
        r2.font.color.rgb = C_TEXT_MAIN

    # ==========================================
    # SLIDE 9: MODULE 6: LIVE ANALYTICAL REPORTS
    # ==========================================
    slide9 = add_slide_base()
    add_slide_header(slide9, "Module 6: Live Analytical Reports & Chair Dashboard", 9)

    rep_cards = [
        ("📊 Summary Metrics", "Live counts for Submissions (6), Accepts (4), Rejects (1), Reviews (16), Reviewers (6).", C_BLUE_LINE),
        ("📈 Track Performance", "Acceptance rates by track: AI/ML (100%), Cloud (50%), Crypto (100%), Software Eng (0%).", C_PURPLE_TEXT),
        ("👥 Reviewer Saturation", "Monitors assigned papers vs capacity; alerts when reviewer approaches max_workload.", C_AMBER_TEXT),
        ("⏳ Pending Review Alerts", "Highlights overdue and pending reviews with reviewer emails for automated follow-up.", C_RED_TEXT)
    ]

    for (title, desc, col), (x, y) in zip(rep_cards, coords):
        c = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = C_WHITE
        c.line.color.rgb = C_BOX_BORDER
        c.line.width = Pt(1)

        st = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(0.12), card_h)
        st.fill.solid()
        st.fill.fore_color.rgb = col
        st.line.fill.background()

        tf = c.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Calibri"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Calibri"
        p.font.size = Pt(12.5)
        p.font.color.rgb = C_TEXT_MAIN

    callout9 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(6.55), Inches(11.5), Inches(0.42))
    callout9.fill.solid()
    callout9.fill.fore_color.rgb = C_GREEN_BG
    callout9.line.color.rgb = C_GREEN_BORDER
    tf_c9 = callout9.text_frame
    tf_c9.margin_top = Inches(0.06)
    p = tf_c9.paragraphs[0]
    p.text = "✔ Seamless Integration: Web dashboard queries SQL Views directly without application-level overhead."
    p.font.name = "Calibri"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = C_GREEN_TEXT
    p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 10: SYSTEMATIC TESTING & VERIFICATION
    # ==========================================
    slide10 = add_slide_base()
    add_slide_header(slide10, "Systematic Testing & Verification Suite (test_suite.py)", 10)

    # Table of tests
    table_shape = slide10.shapes.add_table(6, 4, Inches(0.9), Inches(1.6), Inches(11.533), Inches(4.5))
    table = table_shape.table
    table.columns[0].width = Inches(1.2)
    table.columns[1].width = Inches(3.2)
    table.columns[2].width = Inches(5.3)
    table.columns[3].width = Inches(1.8)

    headers = ["Test #", "Test Name", "Verification Condition", "Status"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY_TITLE
        p = cell.text_frame.paragraphs[0]
        p.font.name = "Calibri"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_WHITE

    test_rows = [
        ("Test 1", "Paper Submission & Versions", "Increments version_num, preserves V1 blind PDF, tags camera-ready copy", "PASSED ✔"),
        ("Test 2", "Conflict of Interest (COI)", "Detects declared institutional/co-author conflict and blocks reviewer assignment", "PASSED ✔"),
        ("Test 3", "Mandatory Review Quota", "Blocks editorial verdict if completed reviews < 3 (verified on Paper #6)", "PASSED ✔"),
        ("Test 4", "Collision-Free Timetable", "Verifies zero overlapping presentation slots for any co-author across all halls", "PASSED ✔"),
        ("Test 5", "Domain Score Constraints", "Verifies all ratings strictly adhere to integer bounds BETWEEN 1 AND 10", "PASSED ✔")
    ]

    for row_idx, data in enumerate(test_rows, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_WHITE if row_idx % 2 == 1 else C_BG_PAGE
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Calibri"
            p.font.size = Pt(11)
            p.font.color.rgb = C_GREEN_TEXT if col_idx == 3 else C_TEXT_MAIN
            if col_idx == 3:
                p.font.bold = True

    callout10 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.42))
    callout10.fill.solid()
    callout10.fill.fore_color.rgb = C_GREEN_BG
    callout10.line.color.rgb = C_GREEN_BORDER
    tf_c10 = callout10.text_frame
    tf_c10.margin_top = Inches(0.06)
    p = tf_c10.paragraphs[0]
    p.text = "✔ Automated Verification: All 5 test cases executed via unittest runner with 100% pass rate in 0.024s."
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_GREEN_TEXT
    p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 11: PROJECT CONCLUSION & RETROSPECTIVE
    # ==========================================
    slide11 = add_slide_base()
    add_slide_header(slide11, "Project Retrospective & Complete Journey", 11)

    box11 = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(1.8), Inches(10.5), Inches(3.2))
    box11.fill.solid()
    box11.fill.fore_color.rgb = C_WHITE
    box11.line.color.rgb = C_BOX_BORDER
    box11.line.width = Pt(1)

    tf11 = box11.text_frame
    tf11.margin_left = Inches(0.4)
    tf11.margin_top = Inches(0.3)

    p = tf11.paragraphs[0]
    p.text = "END-TO-END PROJECT MILESTONES (100% COMPLETE)"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_TITLE
    p.space_after = Pt(12)

    milestones = [
        ("Review 1 (Week 7):", "Conceptual ER design (17 tables), 3NF normalization, and crow's foot diagram."),
        ("Review 2 (Week 10):", "Physical DDL, 50+ rows dataset, 7 analytical views, advanced SQL, indexes & triggers."),
        ("Review 3 (Week 14):", "Flask web app, role-based portals, automated COI & quota gating, 5 unit tests passing.")
    ]
    for m_title, m_desc in milestones:
        p = tf11.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = "✔  " + m_title + " "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = C_GREEN_TEXT
        r2 = p.add_run()
        r2.text = m_desc
        r2.font.size = Pt(12.5)
        r2.font.color.rgb = C_TEXT_MAIN

    ty_box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(5.3), Inches(10.5), Inches(1.2))
    ty_box.fill.solid()
    ty_box.fill.fore_color.rgb = C_BLUE_BG
    ty_box.line.color.rgb = C_BLUE_BORDER
    tf_ty = ty_box.text_frame
    tf_ty.margin_top = Inches(0.18)

    p = tf_ty.paragraphs[0]
    p.text = "Thank You!"
    p.font.name = "Calibri"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_TITLE
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(4)

    p = tf_ty.add_paragraph()
    p.text = "Academic Conference DBMS project is fully implemented, verified, and operational. Ready for Live Demo & Q&A!"
    p.font.name = "Calibri"
    p.font.size = Pt(12.5)
    p.font.color.rgb = C_TEXT_MUTED
    p.alignment = PP_ALIGN.CENTER

    out_file = "Academic_Conference_DBMS_Review3.pptx"
    prs.save(out_file)
    print(f"Review 3 Presentation successfully generated: {out_file}")

if __name__ == "__main__":
    build_review3_presentation()
